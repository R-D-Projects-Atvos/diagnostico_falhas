# -*- coding: utf-8 -*-
"""
Carrega a Matriz de Plantio da planilha para o geodatabase.

A matriz diz, para cada unidade de manejo de solo e faixa de declividade, se
cada periodo do ano e favoravel, aceitavel, restritivo ou favoravel apenas
com irrigacao. Foi validada no Encontro de Campo Grande, na discussao da
Cartilha Agronomica Atvos.

ATENCAO - a classificacao esta na COR DE FUNDO da celula, nao no texto. O
texto das celulas sao os asteriscos, que representam CONDICOES de manejo:

  **   area com cobertura vegetal bem formada
  ***  cobertura bem formada e dessecada com antecedencia
  **** evitar solo argiloso em periodo de frio intenso

Isso importa: "aceitavel **" nao quer dizer que a epoca foi adequada, e sim
que seria adequada SE a cobertura estivesse bem formada. O relatorio precisa
mostrar a condicao junto da classe.

A planilha cobre apenas USL-UEL (polo Sul), 14 unidades de manejo x 3 faixas
de declividade x 17 periodos = 714 combinacoes.

Formato de saida: uma linha por combinacao (formato longo), que e o que o
join por talhao precisa.

Uso: python -u C:\\temp\\carga_matriz_plantio.py

Geotecnologia / Cartografia - Atvos
"""

import datetime
import functools
import os
import arcpy
import openpyxl

print = functools.partial(print, flush=True)
arcpy.env.overwriteOutput = True

# ---------------------------------------------------------------------------
# CONFIGURACAO
# ---------------------------------------------------------------------------

PLANILHA = r"D:\GEO\SOLOS\matriz_plantio.xlsx"
ABA = "Matriz Plantio"

SDE = r"D:\GEO\TALHOES\SQLServer-10-gisdb(atvospublicador).sde"
TB_SAIDA = SDE + r"\ATVOSPUBLICADOR.MATRIZ_PLANTIO"

# cor de fundo -> classe. O vermelho aparece nas duas formas conforme a
# celula tenha sido pintada pela paleta antiga ou pelo seletor de cor.
CLASSES = {
    "FF008000": "Favoravel",
    "FFFFFF00": "Aceitavel",
    "FF92D050": "Favoravel com irrigacao",
    "indexed10": "Restritivo",
    "FFFF0000": "Restritivo",
}

CONDICOES = {
    "**":   "Exige cobertura vegetal bem formada",
    "***":  "Exige cobertura bem formada e dessecada com antecedencia",
    "****": "Evitar solo argiloso em periodo de frio intenso",
}

# coluna da planilha -> periodo. Jan a Mai por quinzena, Jun a Dez por mes.
PERIODOS = [
    ("G", 1, "Jan", "1Q"), ("H", 1, "Jan", "2Q"),
    ("I", 2, "Fev", "1Q"), ("J", 2, "Fev", "2Q"),
    ("K", 3, "Mar", "1Q"), ("L", 3, "Mar", "2Q"),
    ("M", 4, "Abr", "1Q"), ("N", 4, "Abr", "2Q"),
    ("O", 5, "Mai", "1Q"), ("P", 5, "Mai", "2Q"),
    ("Q", 6, "Jun", None), ("R", 7, "Jul", None), ("S", 8, "Ago", None),
    ("T", 9, "Set", None), ("U", 10, "Out", None), ("V", 11, "Nov", None),
    ("W", 12, "Dez", None),
]

LINHA_INICIAL, LINHA_FINAL = 10, 51

CAMPOS = [
    ("POLO", "TEXT", 20),
    ("USINA", "TEXT", 20),
    ("NUM_MANEJO", "SHORT", None),
    ("AGRUP_SOLOS", "TEXT", 254),
    ("FAIXA_DECLIV", "TEXT", 12),
    ("MES_NUM", "SHORT", None),
    ("MES", "TEXT", 5),
    ("QUINZENA", "TEXT", 3),
    ("PERIODO", "TEXT", 12),
    ("CLASSE_EPOCA", "TEXT", 30),
    ("MARCADOR", "TEXT", 6),
    ("CONDICAO", "TEXT", 120),
    ("DATA_CARGA", "DATE", None),
]


def cor(celula):
    f = celula.fill
    if not f or f.patternType != "solid":
        return None
    c = f.fgColor
    if c.type == "rgb" and isinstance(c.rgb, str):
        return c.rgb
    return "%s%s" % (c.type, getattr(c, c.type, ""))


def ler_planilha():
    if not os.path.isfile(PLANILHA):
        raise RuntimeError("planilha nao encontrada: %s" % PLANILHA)

    wb = openpyxl.load_workbook(PLANILHA)
    if ABA not in wb.sheetnames:
        raise RuntimeError("aba '%s' nao encontrada. Abas: %s"
                           % (ABA, ", ".join(wb.sheetnames)))
    ws = wb[ABA]

    linhas = []
    um_atual, solo_atual = None, None
    sem_classe = 0

    for r in range(LINHA_INICIAL, LINHA_FINAL + 1):
        polo = ws.cell(r, 2).value
        usina = ws.cell(r, 3).value
        um = ws.cell(r, 4).value
        solo = ws.cell(r, 5).value
        decl = ws.cell(r, 6).value

        # a unidade e o solo so aparecem na primeira das tres linhas
        if um is not None:
            um_atual = int(um)
            solo_atual = (solo or "").replace("\n", " ").strip()
        if not decl:
            continue
        decl = decl.strip()

        for col, mes_num, mes, quinzena in PERIODOS:
            celula = ws["%s%d" % (col, r)]
            classe = CLASSES.get(cor(celula))
            if classe is None:
                sem_classe += 1
                continue
            marcador = (celula.value or "").strip() if celula.value else None
            periodo = "%s %s" % (mes, quinzena) if quinzena else mes
            linhas.append([
                (polo or "").strip(), (usina or "").strip(), um_atual,
                solo_atual, decl, mes_num, mes, quinzena, periodo,
                classe, marcador, CONDICOES.get(marcador),
            ])

    print("combinacoes lidas: %d" % len(linhas))
    if sem_classe:
        print("  AVISO: %d celulas sem cor reconhecida (ignoradas)" % sem_classe)

    ums = sorted({l[2] for l in linhas})
    faixas = sorted({l[4] for l in linhas})
    print("  unidades de manejo: %s" % ums)
    print("  faixas de declividade: %s" % faixas)

    esperado = len(ums) * len(faixas) * len(PERIODOS)
    if len(linhas) != esperado:
        print("  AVISO: esperava %d combinacoes, li %d" % (esperado, len(linhas)))

    from collections import Counter
    for classe, n in Counter(l[9] for l in linhas).most_common():
        print("    %-26s %4d" % (classe, n))

    return linhas


def criar():
    if arcpy.Exists(TB_SAIDA):
        return
    print("criando a tabela %s..." % TB_SAIDA)
    arcpy.management.CreateTable(SDE, "MATRIZ_PLANTIO")
    for nome, tipo, tam in CAMPOS:
        if tam:
            arcpy.management.AddField(TB_SAIDA, nome, tipo, field_length=tam)
        else:
            arcpy.management.AddField(TB_SAIDA, nome, tipo)
    arcpy.management.AddIndex(TB_SAIDA, ["NUM_MANEJO", "FAIXA_DECLIV"],
                              "IDX_MTZ_UM_DECL")


def carregar():
    linhas = ler_planilha()
    if not linhas:
        raise RuntimeError("planilha sem combinacoes - abortando sem gravar")

    criar()
    print("regravando a tabela...")
    arcpy.management.TruncateTable(TB_SAIDA)

    agora = datetime.datetime.now()
    campos = [c[0] for c in CAMPOS]
    with arcpy.da.InsertCursor(TB_SAIDA, campos) as ins:
        for l in linhas:
            ins.insertRow(l + [agora])

    gravadas = int(arcpy.management.GetCount(TB_SAIDA)[0])
    print("\n=== resultado ===")
    print("  lidas    : %d" % len(linhas))
    print("  gravadas : %d" % gravadas)
    print("\ntabela: %s" % TB_SAIDA)
    print("proximo passo: classificar_epoca_plantio.py")


if __name__ == "__main__":
    carregar()
