# -*- coding: utf-8 -*-
"""
Carrega as estacoes meteorologicas Zeus e o monitoramento diario no SQL Server.

Cria/atualiza:
  ATVOSPUBLICADOR.ESTACOES_ZEUS          feature class de pontos (cadastro)
  ATVOSPUBLICADOR.MONITORAMENTO_ESTACAO  tabela com o diario

Origem provisoria: as duas planilhas exportadas do BigQuery. Quando a
consulta virar tabela no BQ, so muda a funcao de leitura - o esquema e a
logica de carga permanecem.

O cadastro traz o nome no padrao UNIDADE_FAZENDA (ex.: USL_320013), entao a
unidade e o codigo da fazenda sao extraidos dali. Isso permite conferir o
vinculo por cadastro onde ele existe; como so 170 fazendas tem estacao, o
vinculo do relatorio sera por PROXIMIDADE (proximo script).

Substituicao total a cada execucao.

Uso: python -u C:\\temp\\carga_estacoes_zeus.py

Geotecnologia / Cartografia - Atvos
"""

import datetime
import functools
import os
import re
import arcpy
from openpyxl import load_workbook

print = functools.partial(print, flush=True)
arcpy.env.overwriteOutput = True

# ---------------------------------------------------------------------------

XLSX_CADASTRO = r"C:\temp\PICS_-_Estacoes_Zeus.xlsx"
XLSX_DIARIO = r"C:\temp\monitoramento_pic_-_Monitoramento_das_estacoes_Zeus.xlsx"

SDE = r"D:\GEO\TALHOES\SQLServer-10-gisdb(atvospublicador).sde"
DATASET = os.path.join(SDE, "ATVOSPUBLICADOR.AGRICOLA_ATVOS")

FC_ESTACOES = os.path.join(DATASET, "ATVOSPUBLICADOR.ESTACOES_ZEUS")
TB_DIARIO = SDE + r"\ATVOSPUBLICADOR.MONITORAMENTO_ESTACAO"

# Faixa de sanidade da pressao (hPa). Fora disso e leitura espuria do sensor:
# o valor e carregado assim mesmo, mas contado no relatorio da carga.
PRESSAO_MIN_OK = 800.0

CAMPOS_ESTACAO = [
    ("PIC_ID", "LONG", None), ("CLIENT_ID", "LONG", None),
    ("NOME", "TEXT", 40), ("UNIDADE", "TEXT", 10),
    ("COD_FAZENDA", "TEXT", 10), ("LOCALIZACAO", "TEXT", 80),
    ("STATUS", "TEXT", 20), ("LAT", "DOUBLE", None), ("LON", "DOUBLE", None),
    ("RECORDSTAMP", "DATE", None), ("DATA_CARGA", "DATE", None),
]

COLS_DIARIO = [
    ("PIC_ID", "LONG"), ("DIA", "DATE"),
    ("TEMP_MIN", "DOUBLE"), ("TEMP_MEDIA", "DOUBLE"), ("TEMP_MAX", "DOUBLE"),
    ("UMID_MIN", "DOUBLE"), ("UMID_MEDIA", "DOUBLE"), ("UMID_MAX", "DOUBLE"),
    ("PRESSAO_MIN", "DOUBLE"), ("PRESSAO_MEDIA", "DOUBLE"),
    ("PRESSAO_MAX", "DOUBLE"), ("VENTO_INST_MEDIA", "DOUBLE"),
    ("VENTO_MEDIA", "DOUBLE"), ("RAJADA_MAX", "DOUBLE"),
    ("CHUVA_TOTAL", "DOUBLE"), ("IRRADIACAO_MEDIA", "DOUBLE"),
]


def ler_planilha(caminho, aba="Consulta1"):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba]
    it = ws.iter_rows(values_only=True)
    cabecalho = next(it)
    linhas = [l for l in it if l and l[0] is not None]
    wb.close()
    return cabecalho, linhas


def partes_do_nome(nome):
    """USL_320013 -> ('USL', '320013'). Tolera sufixos: 'URC_219053 II'."""
    texto = str(nome or "").strip()
    if "_" not in texto:
        return None, None
    unidade, resto = texto.split("_", 1)
    digitos = re.match(r"\s*(\d+)", resto)
    return unidade.strip().upper(), (digitos.group(1) if digitos else None)


def criar_estacoes():
    if arcpy.Exists(FC_ESTACOES):
        return
    print("criando ESTACOES_ZEUS")
    sr = arcpy.Describe(DATASET).spatialReference
    arcpy.management.CreateFeatureclass(DATASET, "ESTACOES_ZEUS", "POINT",
                                        spatial_reference=sr)
    for nome, tipo, tam in CAMPOS_ESTACAO:
        if tam:
            arcpy.management.AddField(FC_ESTACOES, nome, tipo, field_length=tam)
        else:
            arcpy.management.AddField(FC_ESTACOES, nome, tipo)
    arcpy.management.AddIndex(FC_ESTACOES, ["PIC_ID"], "IDX_EST_PICID")
    arcpy.management.AddIndex(FC_ESTACOES, ["COD_FAZENDA"], "IDX_EST_FAZENDA")


def criar_diario():
    if arcpy.Exists(TB_DIARIO):
        return
    print("criando MONITORAMENTO_ESTACAO")
    arcpy.management.CreateTable(SDE, "MONITORAMENTO_ESTACAO")
    for nome, tipo in COLS_DIARIO:
        arcpy.management.AddField(TB_DIARIO, nome, tipo)
    arcpy.management.AddField(TB_DIARIO, "DATA_CARGA", "DATE")
    arcpy.management.AddIndex(TB_DIARIO, ["PIC_ID"], "IDX_MON_PICID")
    arcpy.management.AddIndex(TB_DIARIO, ["DIA"], "IDX_MON_DIA")


def carregar_estacoes(agora):
    criar_estacoes()
    _, linhas = ler_planilha(XLSX_CADASTRO)
    print("cadastro lido: %d estacoes" % len(linhas))
    if not linhas:
        raise RuntimeError("planilha de cadastro vazia - carga abortada")

    arcpy.management.DeleteFeatures(FC_ESTACOES)

    campos = ["SHAPE@XY"] + [c[0] for c in CAMPOS_ESTACAO]
    sem_faz, por_status = 0, {}
    with arcpy.da.InsertCursor(FC_ESTACOES, campos) as ins:
        for pic, cli, nome, lat, lon, local, status, stamp in linhas:
            unidade, faz = partes_do_nome(nome)
            if not faz:
                sem_faz += 1
            por_status[status] = por_status.get(status, 0) + 1
            ins.insertRow([(lon, lat), pic, cli, nome, unidade, faz,
                           local, status, lat, lon, stamp, agora])

    print("  carregadas: %d" % len(linhas))
    print("  status    : %s" % por_status)
    if sem_faz:
        print("  AVISO: %d estacoes sem codigo de fazenda no nome" % sem_faz)
    return {l[0] for l in linhas}


def carregar_diario(pics_validos, agora):
    criar_diario()
    _, linhas = ler_planilha(XLSX_DIARIO)
    print("monitoramento lido: %d registros" % len(linhas))
    if not linhas:
        raise RuntimeError("planilha de monitoramento vazia - carga abortada")

    arcpy.management.DeleteRows(TB_DIARIO)

    campos = [c[0] for c in COLS_DIARIO] + ["DATA_CARGA"]
    orfaos = pressao_ruim = 0
    dias, chaves = [], set()
    duplicados = 0

    with arcpy.da.InsertCursor(TB_DIARIO, campos) as ins:
        for l in linhas:
            if l[0] not in pics_validos:
                orfaos += 1
            if l[8] is not None and l[8] < PRESSAO_MIN_OK:
                pressao_ruim += 1
            chave = (l[0], l[1])
            if chave in chaves:
                duplicados += 1
            chaves.add(chave)
            dias.append(l[1])
            ins.insertRow(list(l) + [agora])

    print("  carregados: %d" % len(linhas))
    print("  periodo   : %s a %s" % (min(dias).date(), max(dias).date()))
    print("  estacoes   : %d" % len({l[0] for l in linhas}))
    if orfaos:
        print("  AVISO: %d registros de picId fora do cadastro" % orfaos)
    if duplicados:
        print("  AVISO: %d pares estacao+dia repetidos" % duplicados)
    if pressao_ruim:
        print("  nota: %d leituras de pressao abaixo de %.0f hPa "
              "(sensor espurio, carregadas como estao)"
              % (pressao_ruim, PRESSAO_MIN_OK))


def conferir():
    print("\n=== conferencia ===")
    n_est = int(arcpy.management.GetCount(FC_ESTACOES)[0])
    n_mon = int(arcpy.management.GetCount(TB_DIARIO)[0])
    print("  ESTACOES_ZEUS         : %d" % n_est)
    print("  MONITORAMENTO_ESTACAO : %d" % n_mon)

    print("\n  chuva acumulada nos ultimos 30 dias, 5 estacoes:")
    limite = datetime.datetime.now() - datetime.timedelta(days=30)
    acumulado = {}
    with arcpy.da.SearchCursor(TB_DIARIO, ["PIC_ID", "CHUVA_TOTAL", "DIA"]) as cur:
        for pic, chuva, dia in cur:
            if dia and dia >= limite:
                acumulado[pic] = acumulado.get(pic, 0.0) + (chuva or 0.0)
    nomes = {}
    with arcpy.da.SearchCursor(FC_ESTACOES, ["PIC_ID", "NOME"]) as cur:
        for pic, nome in cur:
            nomes[pic] = nome
    for pic in sorted(acumulado, key=lambda p: -acumulado[p])[:5]:
        print("    %-16s %7.1f mm" % (nomes.get(pic, pic), acumulado[pic]))


if __name__ == "__main__":
    momento = datetime.datetime.now()
    pics = carregar_estacoes(momento)
    carregar_diario(pics, momento)
    conferir()
