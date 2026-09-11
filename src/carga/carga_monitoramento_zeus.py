# -*- coding: utf-8 -*-
"""
Atualiza o monitoramento diario das estacoes Zeus (MONITORAMENTO_ESTACAO).

So o diario. O cadastro das estacoes (ESTACOES_ZEUS) muda raramente e fica
com o carga_estacoes_zeus.py, rodado a mao quando precisar.

Origem hoje: a planilha que o Excel atualiza com a consulta de
sql\\monitoramento_zeus.sql, salva em ENTRADAS\\CLIMA na pasta do projeto.
O servidor ainda nao tem acesso ao BigQuery; quando tiver, so a funcao
ler_origem muda - as conferencias e a gravacao continuam as mesmas.

A consulta devolve o historico inteiro ate ontem, e a carga e por
substituicao total. Por isso, ANTES de apagar qualquer coisa, a planilha e
comparada com o que ja esta no banco: se cobrir menos dias, menos estacoes
ou menos linhas, ou se trouxer estacao+dia repetido, a carga para sem tocar
na tabela. Uma planilha exportada pela metade apagaria historico sem erro
nenhum.

Sem --gravar, so simula: le, confere e mostra o que mudaria.

Uso:
  propy -u src\\carga\\carga_monitoramento_zeus.py            (simula)
  propy -u src\\carga\\carga_monitoramento_zeus.py --gravar   (grava)

Codigo de saida: 0 = ok | 1 = nada gravado (conferencia ou erro)

Geotecnologia / Cartografia - Atvos
"""

import datetime
import functools
import os
import sys
from collections import Counter

print = functools.partial(print, flush=True)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------------------
# CONFIGURACAO
# ---------------------------------------------------------------------------

# pasta do projeto no SharePoint, sincronizada pelo OneDrive da conta que roda
# a tarefa agendada - em outra conta este caminho nao existe
PASTA_ENTRADA = (r"C:\Users\joao.fgromboni\OneDrive - Atvos"
                 r"\Geotecnologia-03. Geotecnologia Cartografia - Documentos"
                 r"\Projetos_Cart\DIAGNOSTICO_FALHAS\ENTRADAS\CLIMA")
ABA = "Consulta1"

# a ordem importa: a gravacao e posicional, na ordem de COLS_DIARIO
COLUNAS_ESPERADAS = [
    "picId", "dia", "temp_min", "temp_media", "temp_max",
    "umid_min", "umid_media", "umid_max",
    "pressao_min", "pressao_media", "pressao_max",
    "vento_inst_media", "vento_media", "rajada_max",
    "chuva_total", "irradiacao_media",
]

SDE = r"D:\GEO\TALHOES\SQLServer-10-gisdb(atvospublicador).sde"
TB_DIARIO = SDE + r"\ATVOSPUBLICADOR.MONITORAMENTO_ESTACAO"
FC_ESTACOES = os.path.join(SDE, "ATVOSPUBLICADOR.AGRICOLA_ATVOS",
                           "ATVOSPUBLICADOR.ESTACOES_ZEUS")

# copia da tabela antes de cada gravacao; volta sozinha se a gravacao falhar
GDB_BACKUP = r"D:\GEO\FALHAS\clima_backup.gdb"
TB_BACKUP = "MONITORAMENTO_ANTERIOR"

# planilha que termina mais de tantos dias antes de ontem provavelmente nao foi
# atualizada no Excel: grava assim mesmo, mas avisa no log
DIAS_SEM_ATUALIZAR_ALERTA = 3

# mesma faixa de sanidade do carga_estacoes_zeus: carrega, mas conta no log
PRESSAO_MIN_OK = 800.0


def como_data(v):
    """O Excel devolve datetime; a consulta direta devolvera date."""
    if isinstance(v, datetime.datetime):
        return v.date()
    return v


def arquivo_de_entrada(pasta=PASTA_ENTRADA):
    """A planilha mais recente da pasta. O ~$arquivo que o Excel deixa
    enquanto a planilha esta aberta e ignorado."""
    if not os.path.isdir(pasta):
        raise RuntimeError("pasta de entrada nao encontrada: %s" % pasta)
    candidatos = [os.path.join(pasta, n) for n in os.listdir(pasta)
                  if n.lower().endswith(".xlsx") and not n.startswith("~$")]
    if not candidatos:
        raise RuntimeError("nenhuma planilha .xlsx em %s" % pasta)
    candidatos.sort(key=os.path.getmtime, reverse=True)
    if len(candidatos) > 1:
        print("AVISO: %d planilhas na pasta - usando a mais recente"
              % len(candidatos))
    return candidatos[0]


def ler_origem(caminho):
    """Registros como tuplas na ordem de COLUNAS_ESPERADAS, dia como date."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if ABA not in wb.sheetnames:
            raise RuntimeError("aba '%s' nao encontrada. Abas: %s"
                               % (ABA, ", ".join(wb.sheetnames)))
        it = wb[ABA].iter_rows(values_only=True)
        cabecalho = [("" if c is None else str(c).strip()) for c in next(it)]
        while cabecalho and not cabecalho[-1]:
            cabecalho.pop()
        if cabecalho != COLUNAS_ESPERADAS:
            raise RuntimeError(
                "colunas da planilha diferentes do esperado - a consulta do "
                "Excel mudou?\n  esperado: %s\n  planilha: %s"
                % (COLUNAS_ESPERADAS, cabecalho))
        n = len(COLUNAS_ESPERADAS)
        linhas = []
        for l in it:
            if not l or l[0] is None:
                continue
            l = list(l[:n])
            l[1] = como_data(l[1])
            linhas.append(tuple(l))
    finally:
        wb.close()
    return linhas


def resumo(pares):
    """Tamanho e cobertura de um conjunto de (estacao, dia, ...)."""
    dias = [como_data(p[1]) for p in pares if p[1] is not None]
    repetidos = Counter((p[0], como_data(p[1])) for p in pares)
    return {
        "linhas": len(pares),
        "estacoes": len({p[0] for p in pares}),
        "inicio": min(dias) if dias else None,
        "fim": max(dias) if dias else None,
        "repetidos": sum(1 for n in repetidos.values() if n > 1),
    }


def motivos_para_nao_gravar(novo, atual):
    """Lista vazia = pode gravar.

    Como a carga apaga e regrava, qualquer coisa que a planilha nova tenha a
    menos que o banco seria perdida sem aviso."""
    motivos = []
    if novo["linhas"] == 0:
        motivos.append("a planilha nao tem nenhuma linha")
        return motivos
    if novo["repetidos"]:
        motivos.append("%d pares estacao+dia repetidos na planilha - a chuva "
                       "seria somada em dobro" % novo["repetidos"])
    if atual["linhas"] == 0:
        return motivos
    if novo["inicio"] > atual["inicio"]:
        motivos.append("a planilha comeca em %s e o banco em %s - apagaria os "
                       "dias antigos" % (novo["inicio"], atual["inicio"]))
    if novo["fim"] < atual["fim"]:
        motivos.append("a planilha termina em %s, antes do banco (%s)"
                       % (novo["fim"], atual["fim"]))
    if novo["estacoes"] < atual["estacoes"]:
        motivos.append("a planilha tem %d estacoes e o banco %d"
                       % (novo["estacoes"], atual["estacoes"]))
    if novo["linhas"] < atual["linhas"]:
        motivos.append("a planilha tem %d linhas e o banco %d"
                       % (novo["linhas"], atual["linhas"]))
    return motivos


def gravar(linhas, agora):
    """Copia a tabela atual, apaga e regrava. Se a gravacao falhar no meio,
    devolve a copia antes de sair - a tabela nunca fica pela metade."""
    import arcpy
    from carga_estacoes_zeus import COLS_DIARIO

    arcpy.env.overwriteOutput = True
    if not arcpy.Exists(GDB_BACKUP):
        arcpy.management.CreateFileGDB(os.path.dirname(GDB_BACKUP),
                                       os.path.basename(GDB_BACKUP))
    backup = os.path.join(GDB_BACKUP, TB_BACKUP)
    print("copiando a tabela atual para %s ..." % backup)
    arcpy.conversion.ExportTable(TB_DIARIO, backup)
    print("  copia com %s linhas" % arcpy.management.GetCount(backup)[0])

    campos = [c[0] for c in COLS_DIARIO] + ["DATA_CARGA"]
    print("regravando %d linhas..." % len(linhas))
    arcpy.management.DeleteRows(TB_DIARIO)
    try:
        with arcpy.da.InsertCursor(TB_DIARIO, campos) as ins:
            for l in linhas:
                ins.insertRow(list(l) + [agora])
    except Exception:
        print("ERRO na gravacao - devolvendo a tabela anterior")
        arcpy.management.DeleteRows(TB_DIARIO)
        arcpy.management.Append(backup, TB_DIARIO, "NO_TEST")
        print("  tabela restaurada com %s linhas"
              % arcpy.management.GetCount(TB_DIARIO)[0])
        raise

    gravadas = int(arcpy.management.GetCount(TB_DIARIO)[0])
    if gravadas != len(linhas):
        raise RuntimeError("gravadas %d linhas, esperadas %d"
                           % (gravadas, len(linhas)))
    return gravadas


def main():
    import arcpy

    gravar_de_verdade = "--gravar" in sys.argv
    agora = datetime.datetime.now()
    print("=" * 64)
    print(" MONITORAMENTO ZEUS - %s" % ("GRAVACAO" if gravar_de_verdade
                                        else "SIMULACAO (use --gravar para gravar)"))
    print(" usuario Windows: %s | %s" % (os.environ.get("USERNAME", "?"),
                                          agora.strftime("%d/%m/%Y %H:%M")))
    print("=" * 64)

    caminho = arquivo_de_entrada()
    salva = datetime.datetime.fromtimestamp(os.path.getmtime(caminho))
    print("planilha : %s" % caminho)
    print("salva em : %s" % salva.strftime("%d/%m/%Y %H:%M"))

    linhas = ler_origem(caminho)
    novo = resumo(linhas)
    atual = resumo(list(arcpy.da.SearchCursor(TB_DIARIO, ["PIC_ID", "DIA"])))

    print("\n%-10s %10s %9s %12s %12s" % ("", "linhas", "estacoes", "inicio", "fim"))
    for rotulo, r in (("banco", atual), ("planilha", novo)):
        print("%-10s %10d %9d %12s %12s" % (rotulo, r["linhas"], r["estacoes"],
                                           r["inicio"], r["fim"]))

    validos = {p for (p,) in arcpy.da.SearchCursor(FC_ESTACOES, ["PIC_ID"])}
    orfaos = sum(1 for l in linhas if l[0] not in validos)
    pressao = sum(1 for l in linhas if l[8] is not None and l[8] < PRESSAO_MIN_OK)
    if orfaos:
        print("AVISO: %d registros de estacao fora do cadastro" % orfaos)
    if pressao:
        print("nota: %d leituras de pressao abaixo de %.0f hPa (sensor espurio, "
              "carregadas como estao)" % (pressao, PRESSAO_MIN_OK))

    motivos = motivos_para_nao_gravar(novo, atual)
    if motivos:
        print("\nNADA FOI GRAVADO. Motivo:")
        for m in motivos:
            print("  - %s" % m)
        return 1

    ontem = datetime.date.today() - datetime.timedelta(days=1)
    atraso = (ontem - novo["fim"]).days
    if atraso > DIAS_SEM_ATUALIZAR_ALERTA:
        print("\nALERTA: a planilha termina em %s, %d dias antes de ontem. "
              "Atualize a consulta no Excel e salve em ENTRADAS\\CLIMA."
              % (novo["fim"], atraso))

    novos_dias = sum(1 for l in linhas if atual["fim"] is None or l[1] > atual["fim"])
    print("\nlinhas de dias posteriores ao banco: %d" % novos_dias)

    if not gravar_de_verdade:
        print("SIMULACAO: nada gravado. Com --gravar, a tabela passaria de %d "
              "para %d linhas." % (atual["linhas"], novo["linhas"]))
        return 0

    gravadas = gravar(linhas, agora)
    print("\n[OK] MONITORAMENTO_ESTACAO: %d linhas, %s a %s"
          % (gravadas, novo["inicio"], novo["fim"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
